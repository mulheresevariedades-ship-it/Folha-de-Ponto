"""API do Ponto Digital DIGEP com autenticação JWT, LGPD e despacho de e-mail.

O backend mantém os arquivos fora da pasta pública, separa cada página enviada,
registra o resultado individual do OCR e permite importar servidores por CSV/XLSX.
Quando um mecanismo de OCR não está instalado, o resultado fica explicitamente
marcado como ``OCR indisponível``; nenhum dado é preenchido por demonstração.
CPFs nunca são expostos integralmente em respostas da API.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import mimetypes
import os
import random
import re
import sqlite3
import unicodedata
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import Depends, FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse

import auth
import models
from FolhaPontoBack.processing import process_document
from security import display_cpf, protect_cpf

ROOT = Path(__file__).resolve().parent


def _resolve_data_dir() -> Path:
    """Return a writable directory for uploads and the SQLite database.

    On serverless platforms (Vercel) the project tree is read-only, so the
    runtime data falls back to the platform temp directory when ``data/``
    cannot be written. ``PONTO_DATA_DIR`` forces a specific location.
    """
    env_dir = os.getenv("PONTO_DATA_DIR")
    if env_dir:
        base = Path(env_dir)
        base.mkdir(parents=True, exist_ok=True)
        return base
    candidate = ROOT / "data"
    try:
        candidate.mkdir(parents=True, exist_ok=True)
        probe = candidate / ".write_probe"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink(missing_ok=True)
        return candidate
    except OSError:
        base = Path(os.getenv("TMPDIR") or "/tmp") / "folhaponto"
        base.mkdir(parents=True, exist_ok=True)
        return base


DATA_DIR = _resolve_data_dir()
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = DATA_DIR / "digep.sqlite3"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

IMPORT_PREVIEWS: dict[str, dict[str, Any]] = {}


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def normalize_competency(value: str) -> str:
    month_names = {
        "01": "JANEIRO", "02": "FEVEREIRO", "03": "MARÇO", "04": "ABRIL",
        "05": "MAIO", "06": "JUNHO", "07": "JULHO", "08": "AGOSTO",
        "09": "SETEMBRO", "10": "OUTUBRO", "11": "NOVEMBRO", "12": "DEZEMBRO",
    }
    month, _, year = value.upper().partition("/")
    return f"{month_names.get(month, month)}/{year}" if year else value.upper()


def connect() -> sqlite3.Connection:
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def ensure_column(
    db: sqlite3.Connection, table: str, column: str, definition: str
) -> None:
    columns = {row["name"] for row in db.execute(f"PRAGMA table_info({table})").fetchall()}
    if column not in columns:
        db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db() -> None:
    with connect() as db:
        db.executescript(
            """
            CREATE TABLE IF NOT EXISTS employees (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              name TEXT NOT NULL,
              matricula TEXT NOT NULL UNIQUE,
              cpf TEXT NOT NULL DEFAULT '',
              email TEXT NOT NULL DEFAULT '',
              carga_horaria INTEGER NOT NULL DEFAULT 40,
              acumula_cargo INTEGER NOT NULL DEFAULT 0,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS batches (
              id TEXT PRIMARY KEY,
              original_name TEXT NOT NULL,
              stored_path TEXT NOT NULL,
              sha256 TEXT NOT NULL UNIQUE,
              page_count INTEGER NOT NULL DEFAULT 1,
              mode TEXT NOT NULL,
              created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS timesheets (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              batch_id TEXT,
              employee_id INTEGER,
              name TEXT NOT NULL,
              matricula TEXT,
              competencia TEXT NOT NULL,
              unidade TEXT,
              confidence INTEGER NOT NULL,
              status TEXT NOT NULL,
              note TEXT NOT NULL DEFAULT '',
              reviewed_at TEXT,
              archived_at TEXT,
              source_page INTEGER NOT NULL DEFAULT 1,
              stored_path TEXT,
              sha256 TEXT,
              extracted_text TEXT NOT NULL DEFAULT '',
              processing_mode TEXT NOT NULL DEFAULT 'ocr-indisponivel',
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL,
              FOREIGN KEY(batch_id) REFERENCES batches(id),
              FOREIGN KEY(employee_id) REFERENCES employees(id)
            );
            CREATE TABLE IF NOT EXISTS audit_logs (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              action TEXT NOT NULL,
              entity TEXT NOT NULL,
              entity_id TEXT,
              details TEXT NOT NULL DEFAULT '{}',
              user_id INTEGER,
              created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS users (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              username TEXT NOT NULL UNIQUE,
              password_hash TEXT NOT NULL,
              role TEXT NOT NULL,
              full_name TEXT NOT NULL DEFAULT '',
              active INTEGER NOT NULL DEFAULT 1,
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS email_dispatches (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              timesheet_id INTEGER NOT NULL,
              employee_id INTEGER,
              status TEXT NOT NULL DEFAULT 'pendente',
              authorized_by INTEGER,
              authorized_at TEXT,
              sent_at TEXT,
              simulated_result TEXT NOT NULL DEFAULT '{}',
              error_message TEXT NOT NULL DEFAULT '',
              created_at TEXT NOT NULL,
              updated_at TEXT NOT NULL,
              FOREIGN KEY(timesheet_id) REFERENCES timesheets(id),
              FOREIGN KEY(employee_id) REFERENCES employees(id),
              FOREIGN KEY(authorized_by) REFERENCES users(id)
            );
            CREATE TABLE IF NOT EXISTS revoked_tokens (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              token_hash TEXT NOT NULL UNIQUE,
              user_id INTEGER,
              expires_at TEXT,
              revoked_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_timesheets_competencia ON timesheets(competencia);
            CREATE INDEX IF NOT EXISTS idx_timesheets_matricula ON timesheets(matricula);
            CREATE INDEX IF NOT EXISTS idx_timesheets_status ON timesheets(status);
            CREATE INDEX IF NOT EXISTS idx_dispatches_status ON email_dispatches(status);
            """
        )
        for table, column, definition in (
            ("employees", "cpf", "TEXT NOT NULL DEFAULT ''"),
            ("employees", "updated_at", "TEXT NOT NULL DEFAULT ''"),
            ("timesheets", "stored_path", "TEXT"),
            ("timesheets", "sha256", "TEXT"),
            ("timesheets", "extracted_text", "TEXT NOT NULL DEFAULT ''"),
            ("timesheets", "processing_mode", "TEXT NOT NULL DEFAULT 'ocr-indisponivel'"),
            ("audit_logs", "user_id", "INTEGER"),
        ):
            ensure_column(db, table, column, definition)
        db.execute("DELETE FROM timesheets WHERE batch_id IS NULL")
        timestamp = now()
        duplicated = db.execute(
            """
            SELECT id FROM timesheets
            WHERE status = 'arquivada' AND employee_id IS NOT NULL
              AND id NOT IN (
                SELECT MAX(id) FROM timesheets
                WHERE status = 'arquivada' AND employee_id IS NOT NULL
                GROUP BY employee_id, competencia
              )
            """
        ).fetchall()
        for row in duplicated:
            db.execute(
                "UPDATE timesheets SET status = 'revisao', archived_at = NULL, updated_at = ? WHERE id = ?",
                (timestamp, row["id"]),
            )
        db.execute(
            """
            CREATE UNIQUE INDEX IF NOT EXISTS uq_timesheets_archive_active
            ON timesheets(employee_id, competencia)
            WHERE status = 'arquivada'
            """
        )
        admin_username = os.getenv("ADMIN_USERNAME", "admin@undf.edu.br").strip()
        admin_password = os.getenv("ADMIN_PASSWORD", "").strip()
        if admin_username and admin_password and not db.execute(
            "SELECT id FROM users WHERE username = ?", (admin_username,)
        ).fetchone():
            db.execute(
                """
                INSERT INTO users (username, password_hash, role, full_name, active, created_at, updated_at)
                VALUES (?, ?, 'admin', ?, 1, ?, ?)
                """,
                (admin_username, auth.hash_password(admin_password), "Administrador", timestamp, timestamp),
            )
            audit(db, "create-admin", "user", admin_username, {"username": admin_username})
        db.execute(
            "INSERT INTO audit_logs (action, entity, entity_id, details, user_id, created_at) VALUES (?, ?, ?, ?, ?, ?)",
            ("startup", "system", None, json.dumps({"mode": "real-page-processing"}), None, timestamp),
        )


@asynccontextmanager
async def lifespan(_app: FastAPI):
    try:
        print("[STARTUP] Inicializando banco de dados...")
        init_db()
        print("[STARTUP] Banco de dados inicializado com sucesso!")
        
        print("[STARTUP] Configurando autenticação...")
        auth.configure(
            fetch_user=_fetch_user,
            is_revoked=_is_revoked,
            revoke_token=_revoke_token,
        )
        print("[STARTUP] Autenticação configurada com sucesso!")
        print("[STARTUP] Sistema pronto para operação!")
    except Exception as e:
        print(f"[ERROR] Erro durante startup: {e}")
        import traceback
        traceback.print_exc()
        raise
    
    yield
    
    print("[SHUTDOWN] Sistema encerrando...")


app = FastAPI(title="Ponto Digital DIGEP", version="0.4.0", lifespan=lifespan)


def _fetch_user(user_id: int) -> dict[str, Any] | None:
    with connect() as db:
        row = db.execute(
            "SELECT id, username, role, full_name, active FROM users WHERE id = ?",
            (user_id,),
        ).fetchone()
        return dict(row) if row else None


def _is_revoked(token: str) -> bool:
    digest = hashlib.sha256(token.encode("utf-8")).hexdigest()
    with connect() as db:
        return (
            db.execute("SELECT 1 FROM revoked_tokens WHERE token_hash = ?", (digest,)).fetchone()
            is not None
        )


def _revoke_token(token: str, user_id: int | None, expires_at: Any = None) -> None:
    digest = hashlib.sha256(token.encode("utf-8")).hexdigest()
    with connect() as db:
        db.execute(
            "INSERT INTO revoked_tokens (token_hash, user_id, expires_at, revoked_at) VALUES (?, ?, ?, ?)",
            (
                digest,
                user_id,
                expires_at.isoformat() if isinstance(expires_at, datetime) else expires_at,
                now(),
            ),
        )


def row_to_dict(row: sqlite3.Row) -> dict[str, Any]:
    item = dict(row)
    item["confidence"] = int(item["confidence"])
    item["employee_id"] = item.get("employee_id")
    item.pop("stored_path", None)
    item.pop("sha256", None)
    item["download_url"] = (
        f"/api/timesheets/{item['id']}/download" if row["stored_path"] else None
    )
    item["status_label"] = {
        "reconhecida": "Reconhecida",
        "revisao": "Revisão necessária",
        "baixa_confianca": "Baixa confiança",
        "nao_identificado": "Não identificado",
        "ocr_indisponivel": "OCR indisponível",
        "pendente": "Pendente",
        "arquivada": "Arquivada",
        "rejeitada": "Rejeitada",
    }.get(item["status"], item["status"])
    return item


def audit(
    db: sqlite3.Connection,
    action: str,
    entity: str,
    entity_id: str | int | None,
    details: dict[str, Any],
    user_id: int | None = None,
) -> None:
    db.execute(
        "INSERT INTO audit_logs (action, entity, entity_id, details, user_id, created_at) VALUES (?, ?, ?, ?, ?, ?)",
        (
            action,
            entity,
            str(entity_id) if entity_id is not None else None,
            json.dumps(details),
            user_id,
            now(),
        ),
    )


def upsert_dispatch(
    db: sqlite3.Connection, timesheet_id: int, employee_id: int
) -> None:
    timestamp = now()
    existing = db.execute(
        "SELECT id FROM email_dispatches WHERE timesheet_id = ?", (timesheet_id,)
    ).fetchone()
    if existing:
        db.execute(
            """
            UPDATE email_dispatches
            SET employee_id = ?, status = 'pendente', authorized_by = NULL,
                authorized_at = NULL, sent_at = NULL, simulated_result = '{}',
                error_message = '', updated_at = ?
            WHERE id = ?
            """,
            (employee_id, timestamp, existing["id"]),
        )
    else:
        db.execute(
            """
            INSERT INTO email_dispatches (timesheet_id, employee_id, status, created_at, updated_at)
            VALUES (?, ?, 'pendente', ?, ?)
            """,
            (timesheet_id, employee_id, timestamp, timestamp),
        )


@app.get("/api/health")
def health() -> dict[str, str]:
    return {"status": "ok", "mode": "real", "database": "sqlite"}


@app.post("/api/auth/login")
def login(payload: models.LoginRequest) -> dict[str, Any]:
    with connect() as db:
        row = db.execute(
            "SELECT * FROM users WHERE username = ?", (payload.username,)
        ).fetchone()
        if not row or not row["active"]:
            raise HTTPException(401, "Credenciais inválidas.")
        if not auth.verify_password(payload.password, row["password_hash"]):
            raise HTTPException(401, "Credenciais inválidas.")
        token = auth.create_access_token(
            row["id"], row["username"], row["role"], row["full_name"] or ""
        )
        audit(
            db, "login", "user", row["id"], {"username": row["username"]}, user_id=row["id"]
        )
        return {
            "access_token": token,
            "token_type": "bearer",
            "user": {
                "id": row["id"],
                "username": row["username"],
                "role": row["role"],
                "full_name": row["full_name"],
                "active": bool(row["active"]),
            },
        }


@app.post("/api/auth/logout")
def logout(
    current: dict[str, Any] = Depends(auth.require_roles()),
) -> dict[str, str]:
    auth.revoke_token(current["token"], current["id"], None)
    with connect() as db:
        audit(
            db,
            "logout",
            "user",
            current["id"],
            {"username": current["username"]},
            user_id=current["id"],
        )
    return {"status": "ok"}


@app.get("/api/auth/me")
def me(
    current: dict[str, Any] = Depends(auth.require_roles()),
) -> dict[str, Any]:
    return {key: current[key] for key in ("id", "username", "role", "full_name", "active")}


@app.get("/api/dashboard")
def dashboard(
    competency: str = Query("07/2026"),
    _current: dict[str, Any] = Depends(auth.require_roles()),
) -> dict[str, Any]:
    competency = normalize_competency(competency)
    with connect() as db:
        rows = db.execute(
            "SELECT status, COUNT(*) AS total FROM timesheets WHERE competencia = ? GROUP BY status",
            (competency,),
        ).fetchall()
        counts = {row["status"]: row["total"] for row in rows}
        dispatch_rows = db.execute(
            "SELECT status, COUNT(*) AS total FROM email_dispatches GROUP BY status"
        ).fetchall()
        dispatch_counts = {row["status"]: row["total"] for row in dispatch_rows}
    total = sum(counts.values())
    recognized = counts.get("reconhecida", 0)
    review = counts.get("revisao", 0) + counts.get("pendente", 0)
    low_confidence = counts.get("baixa_confianca", 0)
    unidentified = counts.get("nao_identificado", 0) + counts.get("ocr_indisponivel", 0)
    return {
        "competency": competency,
        "total": total,
        "recognized": recognized,
        "review": review,
        "low_confidence": low_confidence,
        "unidentified": unidentified,
        "ocr_unavailable": counts.get("ocr_indisponivel", 0),
        "archived": counts.get("arquivada", 0),
        "recognition_rate": round((recognized / total) * 100) if total else 0,
        "average_processing_seconds": 0,
        "dispatch_pending": dispatch_counts.get("pendente", 0),
        "dispatch_sent": dispatch_counts.get("enviado", 0),
        "dispatch_error": dispatch_counts.get("erro", 0),
        "demo": False,
        "sample_size": total,
    }


@app.get("/api/timesheets")
def list_timesheets(
    status: str | None = None,
    q: str | None = None,
    competency: str = "07/2026",
    _current: dict[str, Any] = Depends(auth.require_roles()),
) -> list[dict[str, Any]]:
    conditions = ["competencia = ?"]
    params: list[Any] = [normalize_competency(competency)]
    if status and status != "todos":
        conditions.append("status = ?")
        params.append(status)
    if q:
        conditions.append("(name LIKE ? OR matricula LIKE ?)")
        params.extend([f"%{q}%", f"%{q}%"])
    with connect() as db:
        rows = db.execute(
            f"SELECT * FROM timesheets WHERE {' AND '.join(conditions)} ORDER BY id",
            params,
        ).fetchall()
        return [row_to_dict(row) for row in rows]


@app.get("/api/timesheets/{timesheet_id}")
def get_timesheet(
    timesheet_id: int,
    _current: dict[str, Any] = Depends(auth.require_roles()),
) -> dict[str, Any]:
    with connect() as db:
        row = db.execute(
            "SELECT * FROM timesheets WHERE id = ?", (timesheet_id,)
        ).fetchone()
        if not row:
            raise HTTPException(404, "Folha não encontrada.")
        return row_to_dict(row)


@app.post("/api/timesheets/{timesheet_id}/review")
def review_timesheet(
    timesheet_id: int,
    payload: models.ReviewPayload,
    current: dict[str, Any] = Depends(auth.require_roles("admin", "operador")),
) -> dict[str, Any]:
    if payload.employee_id is None and not (payload.note or "").strip():
        raise HTTPException(
            422,
            "Para arquivar é necessário associar um servidor ou informar uma justificativa.",
        )
    with connect() as db:
        current_row = db.execute(
            "SELECT * FROM timesheets WHERE id = ?", (timesheet_id,)
        ).fetchone()
        if not current_row:
            raise HTTPException(404, "Folha não encontrada.")
        employee_id = payload.employee_id
        if employee_id is None and payload.matricula:
            employee = db.execute(
                "SELECT id FROM employees WHERE matricula = ?", (payload.matricula,)
            ).fetchone()
            employee_id = employee["id"] if employee else None
        timestamp = now()
        competency = normalize_competency(payload.competencia)
        if employee_id is not None:
            db.execute(
                """
                UPDATE timesheets
                SET status = 'revisao', archived_at = NULL, updated_at = ?
                WHERE employee_id = ? AND competencia = ? AND status = 'arquivada' AND id <> ?
                """,
                (timestamp, employee_id, competency, timesheet_id),
            )
        db.execute(
            """
            UPDATE timesheets
            SET name = ?, matricula = ?, competencia = ?, employee_id = ?, note = ?,
                status = 'arquivada', reviewed_at = ?, archived_at = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                payload.name,
                payload.matricula,
                competency,
                employee_id,
                payload.note,
                timestamp,
                timestamp,
                timestamp,
                timesheet_id,
            ),
        )
        audit(
            db,
            "archive",
            "timesheet",
            timesheet_id,
            {
                "name": payload.name,
                "employee_id": employee_id,
                "note": payload.note,
                "mode": "manual-review",
            },
            user_id=current["id"],
        )
        if employee_id is not None:
            employee = db.execute(
                "SELECT * FROM employees WHERE id = ?", (employee_id,)
            ).fetchone()
            if employee is not None and int(employee["acumula_cargo"] or 0) == 1:
                upsert_dispatch(db, timesheet_id, employee_id)
        return row_to_dict(
            db.execute("SELECT * FROM timesheets WHERE id = ?", (timesheet_id,)).fetchone()
        )


@app.post("/api/timesheets/{timesheet_id}/pending")
def mark_pending(
    timesheet_id: int,
    note: str = "",
    current: dict[str, Any] = Depends(auth.require_roles("admin", "operador")),
) -> dict[str, Any]:
    with connect() as db:
        if not db.execute(
            "SELECT id FROM timesheets WHERE id = ?", (timesheet_id,)
        ).fetchone():
            raise HTTPException(404, "Folha não encontrada.")
        db.execute(
            "UPDATE timesheets SET status = 'pendente', note = ?, updated_at = ? WHERE id = ?",
            (note[:500], now(), timesheet_id),
        )
        audit(
            db,
            "pending",
            "timesheet",
            timesheet_id,
            {"note": note[:120]},
            user_id=current["id"],
        )
        return row_to_dict(
            db.execute("SELECT * FROM timesheets WHERE id = ?", (timesheet_id,)).fetchone()
        )


@app.post("/api/timesheets/{timesheet_id}/reject")
def reject_timesheet(
    timesheet_id: int,
    current: dict[str, Any] = Depends(auth.require_roles("admin", "operador")),
) -> dict[str, Any]:
    with connect() as db:
        if not db.execute(
            "SELECT id FROM timesheets WHERE id = ?", (timesheet_id,)
        ).fetchone():
            raise HTTPException(404, "Folha não encontrada.")
        db.execute(
            "UPDATE timesheets SET status = 'rejeitada', archived_at = NULL, updated_at = ? WHERE id = ?",
            (now(), timesheet_id),
        )
        audit(
            db,
            "reject",
            "timesheet",
            timesheet_id,
            {},
            user_id=current["id"],
        )
        return row_to_dict(
            db.execute("SELECT * FROM timesheets WHERE id = ?", (timesheet_id,)).fetchone()
        )


def _page_status(extraction: dict[str, Any], employee: sqlite3.Row | None) -> str:
    mode = extraction["mode"]
    if mode == "ocr-indisponivel":
        return "ocr_indisponivel"
    if not extraction["matricula"] or not extraction["nome"] or not extraction["competencia"]:
        return "nao_identificado" if not extraction["matricula"] else "revisao"
    if employee is None:
        return "nao_identificado"
    return "reconhecida" if extraction["confidence"] >= 0.9 else "revisao"


@app.post("/api/batches")
async def create_batch(
    file: UploadFile = File(...),
    current: dict[str, Any] = Depends(auth.require_roles("admin", "operador")),
) -> dict[str, Any]:
    original_name = Path(file.filename or "").name
    suffix = Path(original_name).suffix.lower()
    if suffix not in {".pdf", ".png", ".jpg", ".jpeg"}:
        raise HTTPException(400, "Formato inválido. Envie PDF, PNG, JPG ou JPEG.")
    content = await file.read()
    if not content:
        raise HTTPException(400, "O arquivo enviado está vazio.")
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(413, "O arquivo excede o limite de 20 MB.")
    digest = hashlib.sha256(content).hexdigest()
    with connect() as db:
        if db.execute("SELECT id FROM batches WHERE sha256 = ?", (digest,)).fetchone():
            raise HTTPException(409, "Este arquivo já foi processado.")
        batch_id = f"LOTE-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:5].upper()}"
        target = UPLOAD_DIR / f"{batch_id}{suffix}"
        pages_dir = UPLOAD_DIR / batch_id
        target.write_bytes(content)
        processed = process_document(target, pages_dir)
        page_modes = [page["extraction"]["mode"] for page in processed["page_results"]]
        batch_mode = page_modes[0] if len(set(page_modes)) == 1 and page_modes else "misto"
        db.execute(
            "INSERT INTO batches (id, original_name, stored_path, sha256, page_count, mode, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (batch_id, original_name, str(target.relative_to(DATA_DIR)), digest, processed["pages"], batch_mode, now()),
        )
        created: list[dict[str, Any]] = []
        for page in processed["page_results"]:
            extraction = page["extraction"]
            employee = None
            if extraction["matricula"]:
                employee = db.execute(
                    "SELECT * FROM employees WHERE matricula = ?", (extraction["matricula"],)
                ).fetchone()
            status = _page_status(extraction, employee)
            row_name = extraction["nome"] or "Não identificado"
            competency = (
                normalize_competency(extraction["competencia"])
                if extraction["competencia"]
                else "NÃO RECONHECIDA"
            )
            page_bytes = page["stored_path"].read_bytes()
            page_hash = hashlib.sha256(page_bytes).hexdigest()
            cursor = db.execute(
                """
                INSERT INTO timesheets
                (batch_id, employee_id, name, matricula, competencia, unidade, confidence, status,
                 source_page, stored_path, sha256, extracted_text, processing_mode, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    batch_id,
                    employee["id"] if employee else None,
                    row_name,
                    extraction["matricula"],
                    competency,
                    None,
                    round(extraction["confidence"] * 100),
                    status,
                    page["source_page"],
                    str(page["stored_path"].relative_to(DATA_DIR)),
                    page_hash,
                    extraction["text"],
                    extraction["mode"],
                    now(),
                    now(),
                ),
            )
            created.append(
                {
                    "id": cursor.lastrowid,
                    "source_page": page["source_page"],
                    "status": status,
                    "mode": extraction["mode"],
                    "matricula": extraction["matricula"],
                    "nome": extraction["nome"],
                    "competencia": extraction["competencia"],
                    "confidence": round(extraction["confidence"] * 100),
                }
            )
        audit(
            db,
            "upload",
            "batch",
            batch_id,
            {"pages": processed["pages"], "mode": batch_mode},
            user_id=current["id"],
        )
    return {
        "batch_id": batch_id,
        "pages": processed["pages"],
        "mode": batch_mode,
        "status": "processed",
        "timesheets": created,
    }


@app.get("/api/employees")
def list_employees(
    _current: dict[str, Any] = Depends(auth.require_roles()),
) -> list[dict[str, Any]]:
    with connect() as db:
        rows = db.execute("SELECT * FROM employees ORDER BY name").fetchall()
        return [dict(row, cpf=display_cpf(row["cpf"])) for row in rows]


def _header_key(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


HEADER_MAP = {
    "NOME": "name",
    "MATRICULA": "matricula",
    "CPF": "cpf",
    "E MAIL": "email",
    "EMAIL": "email",
    "CARGA HORARIA": "carga_horaria",
    "ACUMULA CARGO SIM NAO": "acumula_cargo",
    "ACUMULA CARGO": "acumula_cargo",
}
REQUIRED_IMPORT_COLUMNS = {"name", "matricula", "cpf", "email", "carga_horaria", "acumula_cargo"}


def _read_import_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text.splitlines()[0] if text.strip() else "")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        raw_rows = list(reader)
        if not raw_rows:
            raise HTTPException(422, "A planilha está vazia.")
        headers = raw_rows[0]
        values = raw_rows[1:]
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(503, "Suporte XLSX indisponível neste ambiente.") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(422, "A planilha está vazia.")
        headers, values = rows[0], rows[1:]
    else:
        raise HTTPException(400, "Formato inválido. Envie CSV ou XLSX.")
    mapped = [_header_key(header) for header in headers]
    canonical = [HEADER_MAP.get(header) for header in mapped]
    missing = REQUIRED_IMPORT_COLUMNS - {key for key in canonical if key}
    if missing:
        raise HTTPException(422, f"Colunas obrigatórias ausentes: {', '.join(sorted(missing))}.")
    parsed = []
    for values_row in values:
        parsed.append(
            {
                key: (values_row[index] if index < len(values_row) else "")
                for index, key in enumerate(canonical)
                if key
            }
        )
    return parsed


def _validate_import_rows(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    with connect() as db:
        existing = {
            row["matricula"] for row in db.execute("SELECT matricula FROM employees").fetchall()
        }
    seen: set[str] = set()
    valid: list[dict[str, Any]] = []
    invalid: list[dict[str, Any]] = []
    for line_number, raw in enumerate(rows, start=2):
        row = {key: str(value or "").strip() for key, value in raw.items()}
        errors: list[str] = []
        if not row["name"]:
            errors.append("Nome é obrigatório")
        matricula = re.sub(r"\D", "", row["matricula"])
        if not matricula:
            errors.append("Matrícula é obrigatória")
        elif matricula in seen:
            errors.append("Matrícula duplicada na planilha")
        seen.add(matricula)
        cpf = re.sub(r"\D", "", row["cpf"])
        if not cpf:
            errors.append("CPF é obrigatório")
        elif len(cpf) != 11:
            errors.append("CPF deve ter 11 dígitos")
        try:
            carga = int(float(row["carga_horaria"].replace(",", ".")))
            if carga <= 0 or carga > 100:
                raise ValueError
        except ValueError:
            errors.append("Carga Horária deve ser um número entre 1 e 100")
            carga = 0
        acumula_value = _header_key(row["acumula_cargo"])
        if acumula_value not in {"SIM", "NAO"}:
            errors.append("Acumula cargo deve ser Sim ou Não")
        if row["email"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", row["email"]):
            errors.append("E-mail inválido")
        normalized = {
            "name": row["name"],
            "matricula": matricula,
            "cpf": cpf,
            "email": row["email"],
            "carga_horaria": carga,
            "acumula_cargo": int(acumula_value == "SIM"),
            "line": line_number,
            "action": "update" if matricula in existing else "insert",
        }
        if errors:
            invalid.append({"line": line_number, "values": normalized, "errors": errors})
        else:
            valid.append(normalized)
    return valid, invalid


@app.post("/api/employees/import/preview")
async def preview_employee_import(
    file: UploadFile = File(...),
    current: dict[str, Any] = Depends(auth.require_roles("admin", "operador")),
) -> dict[str, Any]:
    filename = Path(file.filename or "").name
    content = await file.read()
    rows = _read_import_rows(filename, content)
    valid, invalid = _validate_import_rows(rows)
    preview_id = uuid.uuid4().hex
    IMPORT_PREVIEWS[preview_id] = {
        "valid": valid,
        "invalid": invalid,
        "filename": filename,
        "created_at": now(),
    }
    return {
        "preview_id": preview_id,
        "filename": filename,
        "total_rows": len(rows),
        "valid_rows": len(valid),
        "invalid_rows": len(invalid),
        "rows": [dict(row, cpf=display_cpf(row["cpf"])) for row in valid],
        "errors": [
            dict(item, values=dict(item["values"], cpf=display_cpf(item["values"]["cpf"])))
            for item in invalid
        ],
    }


@app.post("/api/employees/import/confirm")
def confirm_employee_import(
    payload: models.ImportConfirmPayload,
    current: dict[str, Any] = Depends(auth.require_roles("admin", "operador")),
) -> dict[str, Any]:
    preview = IMPORT_PREVIEWS.pop(payload.preview_id, None)
    if not preview:
        raise HTTPException(404, "Prévia não encontrada ou já confirmada.")
    inserted = updated = 0
    with connect() as db:
        for row in preview["valid"]:
            existing = db.execute(
                "SELECT id FROM employees WHERE matricula = ?", (row["matricula"],)
            ).fetchone()
            stored_cpf = protect_cpf(row["cpf"])
            timestamp = now()
            if existing:
                db.execute(
                    """
                    UPDATE employees
                    SET name = ?, cpf = ?, email = ?, carga_horaria = ?, acumula_cargo = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (row["name"], stored_cpf, row["email"], row["carga_horaria"], row["acumula_cargo"], timestamp, existing["id"]),
                )
                updated += 1
            else:
                db.execute(
                    """
                    INSERT INTO employees (name, matricula, cpf, email, carga_horaria, acumula_cargo, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (row["name"], row["matricula"], stored_cpf, row["email"], row["carga_horaria"], row["acumula_cargo"], timestamp, timestamp),
                )
                inserted += 1
        audit(
            db,
            "import",
            "employees",
            None,
            {"filename": preview["filename"], "inserted": inserted, "updated": updated},
            user_id=current["id"],
        )
    return {
        "status": "confirmed",
        "inserted": inserted,
        "updated": updated,
        "rejected": len(preview["invalid"]),
    }


@app.get("/api/dispatches")
def list_dispatches(
    status: str | None = None,
    _current: dict[str, Any] = Depends(auth.require_roles()),
) -> list[dict[str, Any]]:
    conditions: list[str] = []
    params: list[Any] = []
    if status:
        conditions.append("d.status = ?")
        params.append(status)
    where = f"WHERE {' AND '.join(conditions)}" if conditions else ""
    with connect() as db:
        rows = db.execute(
            f"""
            SELECT d.id, d.timesheet_id, d.employee_id, d.status, d.authorized_by,
                   d.authorized_at, d.sent_at, d.simulated_result, d.error_message,
                   d.created_at, d.updated_at,
                   e.name AS employee_name, e.matricula AS employee_matricula
            FROM email_dispatches d
            LEFT JOIN employees e ON e.id = d.employee_id
            {where}
            ORDER BY d.id
            """,
            params,
        ).fetchall()
        return [dict(row) for row in rows]


@app.post("/api/dispatches/{dispatch_id}/authorize")
def authorize_dispatch(
    dispatch_id: int,
    current: dict[str, Any] = Depends(auth.require_roles("admin")),
) -> dict[str, Any]:
    with connect() as db:
        dispatch = db.execute(
            "SELECT * FROM email_dispatches WHERE id = ?", (dispatch_id,)
        ).fetchone()
        if dispatch is None:
            raise HTTPException(404, "Despacho não encontrado.")
        if dispatch["status"] == "enviado":
            raise HTTPException(409, "Despacho já enviado.")
        db.execute(
            "UPDATE email_dispatches SET authorized_by = ?, authorized_at = ?, updated_at = ? WHERE id = ?",
            (current["id"], now(), now(), dispatch_id),
        )
        audit(
            db,
            "dispatch-authorize",
            "email_dispatch",
            dispatch_id,
            {"timesheet_id": dispatch["timesheet_id"]},
            user_id=current["id"],
        )
        refreshed = db.execute(
            "SELECT * FROM email_dispatches WHERE id = ?", (dispatch_id,)
        ).fetchone()
        return dict(refreshed)


@app.post("/api/dispatches/{dispatch_id}/simulate-send")
def simulate_send(
    dispatch_id: int,
    payload: models.SimulateSendPayload | None = None,
    current: dict[str, Any] = Depends(auth.require_roles("admin", "operador")),
) -> dict[str, Any]:
    error = (
        payload.result == "error"
        if payload is not None
        else random.random() < 0.25
    )
    with connect() as db:
        dispatch = db.execute(
            "SELECT * FROM email_dispatches WHERE id = ?", (dispatch_id,)
        ).fetchone()
        if dispatch is None:
            raise HTTPException(404, "Despacho não encontrado.")
        if dispatch["status"] == "enviado":
            raise HTTPException(409, "Despacho já enviado.")
        timestamp = now()
        if error:
            message = "Falha simulada no envio para o servidor associado."
            simulated = json.dumps({"status": "erro", "simulated": True, "message": message})
            db.execute(
                """
                UPDATE email_dispatches
                SET status = 'erro', sent_at = NULL, simulated_result = ?, error_message = ?, updated_at = ?
                WHERE id = ?
                """,
                (simulated, message, timestamp, dispatch_id),
            )
        else:
            simulated = json.dumps(
                {"status": "enviado", "simulated": True, "delivered": True}
            )
            db.execute(
                """
                UPDATE email_dispatches
                SET status = 'enviado', sent_at = ?, simulated_result = ?, error_message = '', updated_at = ?
                WHERE id = ?
                """,
                (timestamp, simulated, timestamp, dispatch_id),
            )
        audit(
            db,
            "dispatch-simulate",
            "email_dispatch",
            dispatch_id,
            {
                "timesheet_id": dispatch["timesheet_id"],
                "result": "erro" if error else "enviado",
            },
            user_id=current["id"],
        )
        refreshed = db.execute(
            "SELECT * FROM email_dispatches WHERE id = ?", (dispatch_id,)
        ).fetchone()
        return dict(refreshed)


@app.post("/api/dispatches/{dispatch_id}/resend")
def resend_dispatch(
    dispatch_id: int,
    current: dict[str, Any] = Depends(auth.require_roles("admin", "operador")),
) -> dict[str, Any]:
    with connect() as db:
        dispatch = db.execute(
            "SELECT * FROM email_dispatches WHERE id = ?", (dispatch_id,)
        ).fetchone()
        if dispatch is None:
            raise HTTPException(404, "Despacho não encontrado.")
        db.execute(
            """
            UPDATE email_dispatches
            SET status = 'pendente', authorized_by = NULL, authorized_at = NULL,
                sent_at = NULL, simulated_result = '{}', error_message = '', updated_at = ?
            WHERE id = ?
            """,
            (now(), dispatch_id),
        )
        audit(
            db,
            "dispatch-resend",
            "email_dispatch",
            dispatch_id,
            {"timesheet_id": dispatch["timesheet_id"]},
            user_id=current["id"],
        )
        refreshed = db.execute(
            "SELECT * FROM email_dispatches WHERE id = ?", (dispatch_id,)
        ).fetchone()
        return dict(refreshed)


@app.get("/api/timesheets/{timesheet_id}/download")
def download_timesheet(
    timesheet_id: int,
    _current: dict[str, Any] = Depends(auth.require_roles()),
) -> FileResponse:
    with connect() as db:
        row = db.execute(
            "SELECT stored_path FROM timesheets WHERE id = ?", (timesheet_id,)
        ).fetchone()
    if not row or not row["stored_path"]:
        raise HTTPException(404, "Arquivo individual não encontrado.")
    path = (DATA_DIR / row["stored_path"]).resolve()
    if DATA_DIR.resolve() not in path.parents or not path.is_file():
        raise HTTPException(404, "Arquivo individual não encontrado.")
    media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return FileResponse(path, filename=path.name, media_type=media_type)


@app.get("/")
def index() -> FileResponse:
    """Servir página HTML principal."""
    return FileResponse(ROOT / "index.html")


@app.get("/styles.css", include_in_schema=False)
def styles() -> FileResponse:
    """Servir a folha de estilos referenciada pelo HTML principal."""
    return FileResponse(ROOT / "styles.css", media_type="text/css")


@app.get("/app.js", include_in_schema=False)
def javascript() -> FileResponse:
    """Servir o JavaScript referenciado pelo HTML principal."""
    return FileResponse(ROOT / "app.js", media_type="application/javascript")